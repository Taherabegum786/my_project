"""Generates Course_Plan_CS2101.docx from the same data as the PDF."""
import runpy, sys, types
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, Mm

# reuse UNITS / modes from the PDF script without building the PDF
src = open(__file__.replace("build_course_plan_docx.py", "build_course_plan.py")).read()
ns = {"__file__": "x"}
exec(src.split("GRID = [")[0], ns)
UNITS, COURSE, PROGRAMME, FACULTY = ns["UNITS"], ns["COURSE"], ns["PROGRAMME"], ns["FACULTY"]
exec("modes = [" + src.split("modes = [")[1].split("]\n")[0] + "]", ns)
MODES = ns["modes"]
OUT = __file__.replace("build_course_plan_docx.py", "Course_Plan_CS2101.docx")
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
for side in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
    setattr(sec, side, Mm(12))

normal = doc.styles["Normal"]
normal.font.name, normal.font.size = FONT, Pt(9)
normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
normal.paragraph_format.space_after = Pt(0)
normal.paragraph_format.space_before = Pt(0)


def para(text="", size=9, bold=False, align=None, before=0, after=0, underline=False, keep=False):
    p = doc.add_paragraph()
    if text:
        r = p.add_run(text)
        r.font.size, r.bold, r.underline = Pt(size), bold, underline
    if align is not None:
        p.alignment = align
    p.paragraph_format.space_before, p.paragraph_format.space_after = Pt(before), Pt(after)
    p.paragraph_format.keep_with_next = keep
    return p


def shade(cell, fill="E6E6E6"):
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), fill)
    cell._tc.get_or_add_tcPr().append(shd)


def repeat_header(row):
    trPr = row._tr.get_or_add_trPr()
    el = OxmlElement("w:tblHeader"); el.set(qn("w:val"), "true"); trPr.append(el)


def cant_split(row):
    el = OxmlElement("w:cantSplit"); el.set(qn("w:val"), "true"); row._tr.get_or_add_trPr().append(el)


def table(header, rows, widths_mm, centre_cols=(), bold_last=False, keep_all=False):
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    n = len(rows) + 1
    for ri, row_vals in enumerate([header] + list(rows)):
        row = t.rows[0] if ri == 0 else t.add_row()
        cant_split(row)
        for i, (val, cell) in enumerate(zip(row_vals, row.cells)):
            cell.width = Mm(widths_mm[i])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            p = cell.paragraphs[0]
            p.paragraph_format.keep_with_next = ri == 0 or (keep_all and ri < n - 1)
            r = p.add_run(str(val))
            r.font.size = Pt(9)
            r.bold = ri == 0 or (bold_last and ri == n - 1 and i > 0)
            if i in centre_cols:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    mar = OxmlElement("w:tblCellMar")
    for side in ("top", "bottom"):
        el = OxmlElement(f"w:{side}"); el.set(qn("w:w"), "20"); el.set(qn("w:type"), "dxa"); mar.append(el)
    t._tbl.tblPr.append(mar)
    repeat_header(t.rows[0])
    for c in t.rows[0].cells:
        shade(c)
    # column widths at grid level too (Word honours cell widths, LibreOffice/Google the grid)
    for i, col in enumerate(t.columns):
        col.width = Mm(widths_mm[i])
    return t


# ---- cover page
para(before=200)
for text, size, bold, gap in [(PROGRAMME, 16, True, 8), ("COURSE PLAN", 18, True, 8), (COURSE, 22, True, 60),
                              ("Prepared by", 12, False, 4), (FACULTY, 18, True, 4),
                              ("Department of CSSE", 12, False, 2), ("Andhra University", 12, True, 0)]:
    para(text, size, bold, WD_ALIGN_PARAGRAPH.CENTER, after=gap)
doc.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

# ---- unit-wise plan
para(PROGRAMME, 13, True, WD_ALIGN_PARAGRAPH.CENTER, after=4)
para("UNIT-WISE COURSE PLAN", 15, True, WD_ALIGN_PARAGRAPH.CENTER, after=6, underline=True)
para(COURSE, 11.5, True, WD_ALIGN_PARAGRAPH.CENTER, after=6, underline=True)
p = para(after=2)
for txt, b in [("Total Hours: ", True), ("60 (15 weeks × 4 hours)", False)]:
    r = p.add_run(txt); r.bold = b; r.font.size = Pt(9.5)
p = para(after=2)
for txt, b, it in [("Text Book: ", True, False), ("Kenneth H. Rosen, ", False, False),
                   ("Discrete Mathematics & Its Applications with Combinatorics and Graph Theory", False, True),
                   (", Tata McGraw-Hill", False, False)]:
    r = p.add_run(txt); r.bold, r.italic = b, it; r.font.size = Pt(9.5)

dist = []
for num, name, topics in UNITS:
    hrs = sum(t[1] for t in topics)
    dist.append((num, name, hrs))
    para(f"UNIT {num}: {name} — {hrs} Hours", 11.5, True, before=6, after=3, keep=True)
    table(["S.No.", "Topic", "Hrs", "Mode of Teaching", "Learning Activity"],
          [(i, *t) for i, t in enumerate(topics, 1)], [12, 73, 11, 35, 55], centre_cols=(0, 2))
total = sum(d[2] for d in dist)

para("Overall Hour Distribution", 11.5, True, before=6, after=3, keep=True)
table(["Unit", "Unit Title", "Hours", "Percentage"],
      [(n, t, h, f"{h / total * 100:.2f}%") for n, t, h in dist] + [("", "GRAND TOTAL", total, "100%")],
      [16, 110, 30, 30], centre_cols=(0, 2, 3), bold_last=True, keep_all=True)

para("General Modes of Teaching", 11.5, True, before=6, after=3, keep=True)
for m in MODES:
    p = doc.add_paragraph(style="List Bullet")
    r = p.add_run(m); r.font.size = Pt(9.5)
    p.paragraph_format.space_after = Pt(0)

doc.core_properties.title = "Discrete Mathematical Structures – Course Plan"
doc.core_properties.author = FACULTY
doc.save(OUT)
print("wrote", OUT)
